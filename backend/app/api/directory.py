import csv
import io

from fastapi import APIRouter, Depends, File, Query, Request, UploadFile
from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from app.api.util import PageParams, like, paginate
from app.core.db import get_db
from app.core.deps import Auth, client_ip, require_role
from app.core.errors import ApiError
from app.models import Department, Employee, Role
from app.schemas.common import (
    DepartmentIn,
    DepartmentOut,
    DepartmentUpdate,
    EmployeeIn,
    EmployeeOut,
    EmployeeUpdate,
    ImportResult,
    ImportRow,
    Page,
)
from app.services.audit import audit

router = APIRouter(tags=["directory"])
_email = TypeAdapter(EmailStr)
MAX_IMPORT_BYTES, MAX_IMPORT_ROWS = 5 * 1024 * 1024, 5000


def _dept_out(db: Session, d: Department) -> DepartmentOut:
    n = db.scalar(
        select(func.count(Employee.id)).where(
            Employee.department_id == d.id, Employee.is_active.is_(True)
        )
    )
    return DepartmentOut(id=d.id, name=d.name, is_active=d.is_active, employee_count=n)


def _emp_out(e: Employee) -> EmployeeOut:
    out = EmployeeOut.model_validate(e)
    out.department_name = e.department.name if e.department else None
    return out


# ---------------------------------------------------------------- departments
@router.get("/departments", response_model=list[DepartmentOut])
def list_departments(
    include_inactive: bool = False, _: Auth = Depends(require_role(Role.NOTICE_EDITOR)),
    db: Session = Depends(get_db),
):  # fmt: skip
    counts = dict(
        db.execute(
            select(Employee.department_id, func.count(Employee.id))
            .where(Employee.is_active.is_(True))
            .group_by(Employee.department_id)
        ).all()
    )
    q = select(Department).order_by(func.lower(Department.name))
    if not include_inactive:
        q = q.where(Department.is_active.is_(True))
    return [
        DepartmentOut(
            id=d.id, name=d.name, is_active=d.is_active, employee_count=counts.get(d.id, 0)
        )
        for d in db.scalars(q)
    ]


def _dept_name_taken(db: Session, name: str, exclude_id: int | None = None) -> bool:
    q = select(Department.id).where(func.lower(Department.name) == name.lower())
    if exclude_id:
        q = q.where(Department.id != exclude_id)
    return db.scalar(q) is not None


@router.post("/departments", response_model=DepartmentOut, status_code=201)
def create_department(
    body: DepartmentIn, request: Request, auth: Auth = Depends(require_role(Role.ADMIN)),
    db: Session = Depends(get_db),
):  # fmt: skip
    if _dept_name_taken(db, body.name):
        raise ApiError(409, "duplicate", "A department with this name already exists")
    d = Department(name=body.name)
    db.add(d)
    db.flush()
    audit(db, "department.create", auth.user.id, "department", d.id, ip=client_ip(request))
    db.commit()
    return _dept_out(db, d)


@router.patch("/departments/{dept_id}", response_model=DepartmentOut)
def update_department(
    dept_id: int, body: DepartmentUpdate, request: Request,
    auth: Auth = Depends(require_role(Role.ADMIN)), db: Session = Depends(get_db),
):  # fmt: skip
    d = db.get(Department, dept_id)
    if not d:
        raise ApiError(404, "not_found", "Department not found")
    data = body.model_dump(exclude_unset=True)
    if "name" in data:
        data["name"] = data["name"].strip()
        if _dept_name_taken(db, data["name"], d.id):
            raise ApiError(409, "duplicate", "A department with this name already exists")
    for k, v in data.items():
        setattr(d, k, v)
    audit(db, "department.update", auth.user.id, "department", d.id, data, client_ip(request))
    db.commit()
    return _dept_out(db, d)


# ------------------------------------------------------------------ employees
@router.get("/employees", response_model=Page[EmployeeOut])
def list_employees(
    search: str | None = None, department_id: int | None = None, active: bool | None = None,
    p: PageParams = Depends(), _: Auth = Depends(require_role(Role.NOTICE_EDITOR)),
    db: Session = Depends(get_db),
):  # fmt: skip
    q = (
        select(Employee)
        .options(joinedload(Employee.department))
        .order_by(Employee.name, Employee.id)
    )
    if search:
        t = like(search.strip())
        q = q.where(or_(Employee.name.ilike(t, escape="\\"), Employee.email.ilike(t, escape="\\")))
    if department_id:
        q = q.where(Employee.department_id == department_id)
    if active is not None:
        q = q.where(Employee.is_active.is_(active))
    total, rows = paginate(db, q, p)
    return Page(items=[_emp_out(e) for e in rows.scalars().unique()], total=total,
                page=p.page, page_size=p.page_size)  # fmt: skip


def _check_department(db: Session, dept_id: int | None):
    if dept_id is not None and not db.get(Department, dept_id):
        raise ApiError(422, "invalid_department", "Department does not exist")


@router.post("/employees", response_model=EmployeeOut, status_code=201)
def create_employee(
    body: EmployeeIn, request: Request, auth: Auth = Depends(require_role(Role.ADMIN)),
    db: Session = Depends(get_db),
):  # fmt: skip
    _check_department(db, body.department_id)
    email = body.email.lower()
    if db.scalar(select(Employee.id).where(Employee.email == email)):
        raise ApiError(409, "duplicate_email", "An employee with this email already exists")
    e = Employee(name=body.name.strip(), email=email, department_id=body.department_id,
                 job_title=body.job_title)  # fmt: skip
    db.add(e)
    db.flush()
    audit(db, "employee.create", auth.user.id, "employee", e.id, ip=client_ip(request))
    db.commit()
    db.refresh(e)
    return _emp_out(e)


@router.patch("/employees/{emp_id}", response_model=EmployeeOut)
def update_employee(
    emp_id: int, body: EmployeeUpdate, request: Request,
    auth: Auth = Depends(require_role(Role.ADMIN)), db: Session = Depends(get_db),
):  # fmt: skip
    e = db.get(Employee, emp_id)
    if not e:
        raise ApiError(404, "not_found", "Employee not found")
    data = body.model_dump(exclude_unset=True)
    if "department_id" in data:
        _check_department(db, data["department_id"])
    if "email" in data:
        data["email"] = data["email"].lower()
        if db.scalar(
            select(Employee.id).where(Employee.email == data["email"], Employee.id != e.id)
        ):
            raise ApiError(409, "duplicate_email", "An employee with this email already exists")
    for k, v in data.items():
        setattr(e, k, v)
    audit(db, "employee.update", auth.user.id, "employee", e.id, {"fields": sorted(data)},
          client_ip(request))  # fmt: skip
    db.commit()
    db.refresh(e)
    return _emp_out(e)


@router.delete("/employees/{emp_id}", response_model=EmployeeOut)
def deactivate_employee(
    emp_id: int, request: Request, auth: Auth = Depends(require_role(Role.ADMIN)),
    db: Session = Depends(get_db),
):  # fmt: skip
    """Soft delete: history (past notice recipients) is preserved."""
    e = db.get(Employee, emp_id)
    if not e:
        raise ApiError(404, "not_found", "Employee not found")
    e.is_active = False
    audit(db, "employee.deactivate", auth.user.id, "employee", e.id, ip=client_ip(request))
    db.commit()
    return _emp_out(e)


# --------------------------------------------------------------------- import
_ALIASES = {
    "name": {"name", "employee name", "full name", "employee"},
    "email": {"email", "email address", "e-mail", "mail"},
    "department": {"department", "dept"},
    "job_title": {"job title", "title", "designation", "position"},
}


def _read_rows(filename: str, data: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ApiError(422, "bad_file", "CSV must be UTF-8 encoded") from None
        raw = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            raw = [
                ["" if c is None else str(c) for c in r]
                for r in wb.active.iter_rows(values_only=True)
            ]
        except Exception:  # noqa: BLE001
            raise ApiError(422, "bad_file", "Could not read the Excel file") from None
    else:
        raise ApiError(422, "bad_file", "Upload a .csv or .xlsx file")
    if not raw:
        raise ApiError(422, "bad_file", "The file is empty")
    header = [h.strip().lower() for h in raw[0]]
    cols = {}
    for field, names in _ALIASES.items():
        for i, h in enumerate(header):
            if h in names:
                cols[field] = i
                break
    if "name" not in cols or "email" not in cols:
        raise ApiError(422, "bad_file", "File needs 'Name' and 'Email' columns")
    out = []
    for r in raw[1:]:
        if not any(c.strip() for c in r):
            continue
        out.append({f: (r[i].strip() if i < len(r) else "") for f, i in cols.items()})
    if len(out) > MAX_IMPORT_ROWS:
        raise ApiError(422, "bad_file", f"Maximum {MAX_IMPORT_ROWS} rows per import")
    return out


@router.post("/employees/import", response_model=ImportResult)
async def import_employees(
    request: Request, file: UploadFile = File(...), commit: bool = Query(False),
    auth: Auth = Depends(require_role(Role.ADMIN)), db: Session = Depends(get_db),
):  # fmt: skip
    """commit=false previews; commit=true applies. Existing emails are skipped, never overwritten."""
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise ApiError(413, "bad_file", "Import file exceeds 5 MB")
    rows = _read_rows(file.filename or "", data)
    depts = {d.name.lower(): d for d in db.scalars(select(Department))}
    existing = set(db.scalars(select(Employee.email)))
    seen: set[str] = set()
    results: list[ImportRow] = []
    to_create: list[tuple[str, str, str | None, str | None]] = []
    new_depts: dict[str, str] = {}
    for i, r in enumerate(rows, start=2):
        row = ImportRow(row=i, name=r["name"], email=r.get("email", ""),
                        department=r.get("department") or None,
                        job_title=r.get("job_title") or None, status="ok")  # fmt: skip
        try:
            email = str(_email.validate_python(row.email)).lower()
            row.email = email
        except ValidationError:
            row.status, row.message = "invalid", "Invalid email address"
            results.append(row)
            continue
        if not row.name:
            row.status, row.message = "invalid", "Name is required"
            results.append(row)
            continue
        if email in seen:
            row.status, row.message = "duplicate_in_file", "Duplicate of an earlier row"
        elif email in existing:
            row.status, row.message = "exists", "Employee already exists (not changed)"
        else:
            if row.department and row.department.lower() not in depts:
                new_depts.setdefault(row.department.lower(), row.department)
                row.new_department = True
            to_create.append((row.name, email, row.department, row.job_title))
        seen.add(email)
        results.append(row)
    if commit:
        try:
            for key, name in new_depts.items():
                d = Department(name=name)
                db.add(d)
                db.flush()
                depts[key] = d
            for name, email, dept, title in to_create:
                db.add(Employee(name=name, email=email, job_title=title,
                                department_id=depts[dept.lower()].id if dept else None))  # fmt: skip
            audit(db, "employee.import", auth.user.id, detail={"created": len(to_create),
                  "file": file.filename}, ip=client_ip(request))  # fmt: skip
            db.commit()
        except IntegrityError:
            db.rollback()
            raise ApiError(
                409, "import_conflict", "Data changed during import; re-run preview"
            ) from None
    return ImportResult(
        committed=commit, total=len(results), ok=len(to_create),
        skipped=sum(r.status in ("exists", "duplicate_in_file") for r in results),
        errors=sum(r.status == "invalid" for r in results), rows=results,
    )  # fmt: skip
